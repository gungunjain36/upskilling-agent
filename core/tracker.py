from datetime import datetime
from pathlib import Path
from typing import Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from config import settings

SHEETS = {
    "roadmap": "Roadmap",
    "sessions": "Sessions",
    "challenges": "Challenges",
    "concepts": "Concepts",
    "profile": "Profile",
}


def _get_or_create_workbook() -> openpyxl.Workbook:
    path = settings.tracker_path
    path.parent.mkdir(parents=True, exist_ok=True)
    if path.exists():
        return openpyxl.load_workbook(path)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _init_sheets(wb)
    wb.save(path)
    return wb


def _save_workbook(wb: openpyxl.Workbook):
    wb.save(settings.tracker_path)


def _header_style(cell):
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill(fill_type="solid", fgColor="2D6A4F")
    cell.alignment = Alignment(horizontal="center")


def _init_sheets(wb: openpyxl.Workbook):
    _init_roadmap_sheet(wb)
    _init_sessions_sheet(wb)
    _init_challenges_sheet(wb)
    _init_concepts_sheet(wb)
    _init_profile_sheet(wb)


def _init_roadmap_sheet(wb):
    ws = wb.create_sheet(SHEETS["roadmap"])
    headers = ["Topic", "Domain", "Level Target", "Current Level", "Status", "Priority", "Notes", "Last Updated"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        _header_style(cell)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["G"].width = 40


def _init_sessions_sheet(wb):
    ws = wb.create_sheet(SHEETS["sessions"])
    headers = ["Timestamp", "Domain", "Topic", "Type", "Summary", "Duration (min)"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        _header_style(cell)
    ws.column_dimensions["E"].width = 50


def _init_challenges_sheet(wb):
    ws = wb.create_sheet(SHEETS["challenges"])
    headers = ["Timestamp", "Domain", "Topic", "Challenge", "User Response", "Evaluation", "Score (1-10)", "Status"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        _header_style(cell)
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["E"].width = 50


def _init_concepts_sheet(wb):
    ws = wb.create_sheet(SHEETS["concepts"])
    headers = ["Timestamp", "Domain", "Topic", "Concept Title", "Concept Summary", "Delivered"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        _header_style(cell)
    ws.column_dimensions["E"].width = 60


def _init_profile_sheet(wb):
    ws = wb.create_sheet(SHEETS["profile"])
    headers = ["Key", "Value", "Updated At"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=i, value=h)
        _header_style(cell)
    ws.column_dimensions["A"].width = 25
    ws.column_dimensions["B"].width = 50


def init_tracker():
    """Ensure the tracker file exists with all sheets."""
    wb = _get_or_create_workbook()
    existing = set(wb.sheetnames)
    for sheet_name in SHEETS.values():
        if sheet_name not in existing:
            _init_sheets(wb)
            break
    _save_workbook(wb)


def log_session(domain: str, topic: str, session_type: str, summary: str, duration: int = 0):
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["sessions"]]
    ws.append([datetime.now().isoformat(), domain, topic, session_type, summary, duration])
    _save_workbook(wb)


def log_concept(domain: str, topic: str, title: str, summary: str):
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["concepts"]]
    ws.append([datetime.now().isoformat(), domain, topic, title, summary, True])
    _save_workbook(wb)


def log_challenge(domain: str, topic: str, challenge: str, user_response: str = "", evaluation: str = "", score: int = 0, status: str = "pending"):
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["challenges"]]
    ws.append([datetime.now().isoformat(), domain, topic, challenge, user_response, evaluation, score, status])
    _save_workbook(wb)


def update_roadmap(topic: str, domain: str, level_target: int, current_level: int, status: str, priority: int = 1, notes: str = ""):
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["roadmap"]]

    # Check if topic exists and update
    for row in ws.iter_rows(min_row=2):
        if row[0].value == topic and row[1].value == domain:
            row[3].value = current_level
            row[4].value = status
            row[7].value = datetime.now().isoformat()
            _save_workbook(wb)
            return

    # Add new row
    ws.append([topic, domain, level_target, current_level, status, priority, notes, datetime.now().isoformat()])
    _save_workbook(wb)


def update_profile_in_tracker(key: str, value: str):
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["profile"]]
    now = datetime.now().isoformat()

    for row in ws.iter_rows(min_row=2):
        if row[0].value == key:
            row[1].value = value
            row[2].value = now
            _save_workbook(wb)
            return

    ws.append([key, value, now])
    _save_workbook(wb)


def get_roadmap_summary() -> list[dict]:
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["roadmap"]]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            items.append({
                "topic": row[0], "domain": row[1], "level_target": row[2],
                "current_level": row[3], "status": row[4], "priority": row[5],
            })
    return items


def get_challenges_for_topic(topic: str) -> list[dict]:
    wb = _get_or_create_workbook()
    ws = wb[SHEETS["challenges"]]
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[2] == topic:
            items.append({
                "timestamp": row[0], "domain": row[1], "topic": row[2],
                "challenge": row[3], "score": row[6], "status": row[7],
            })
    return items
